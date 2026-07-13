import csv
import io
import json
import os
import urllib.error
import urllib.request
from datetime import datetime, timedelta

from fastapi import APIRouter, Body, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from sqlalchemy import func, or_, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from sqlalchemy.orm import selectinload
from starlette.concurrency import run_in_threadpool

import models
from config import settings
from core.dependencies import get_current_user, get_db, templates
from core.finance import credit_user_balance
from core.fraud import evaluate_admin_pc_override, evaluate_admin_top_up
from core.integrations import check_integration, touch_integration_status, validate_integration_url
from core.roles import CLUB_ADMIN_ROLES, ROLE_PENDING_OWNER, ROLE_SUPERADMIN, STAFF_ROLES
from core.services.booking_service import BookingService
from core.services.exceptions import ServiceError
from core.services.order_service import OrderService
from core.tenancy import can_manage_club, can_manage_order, can_manage_pc, manageable_club_ids
from core.ws import issue_pc_token

router = APIRouter(prefix="/admin", tags=["web_admin"])

VALID_PC_STATUSES = {"free", "busy", "reserved"}
VALID_PC_CATEGORIES = {"Standard", "VIP"}
MAX_PACKAGE_DURATION_MINUTES = 24 * 60


async def create_notification(db: AsyncSession, kind: str, title: str, message: str, club_id=None):
    db.add(models.Notification(kind=kind, title=title, message=message, club_id=club_id))


async def get_owner_clubs(db: AsyncSession, owner: models.User):
    res = await db.execute(
        select(models.Club)
        .filter(models.Club.id.in_(await manageable_club_ids(db, owner)))
        .order_by(models.Club.name)
    )
    return res.scalars().all()


async def owner_club_ids(db: AsyncSession, owner: models.User):
    return await manageable_club_ids(db, owner)


async def ensure_owner_club(db: AsyncSession, owner: models.User, club_id: int):
    return await can_manage_club(db, owner, club_id)


def normalize_text(value: str | None) -> str:
    return " ".join((value or "").strip().split())


def amount(value: int | None) -> int:
    return int(value or 0)


def package_total_minutes(paid_minutes: int, bonus_minutes: int) -> int:
    return paid_minutes + bonus_minutes


def build_finance_summary(orders, bookings, expenses):
    order_revenue = sum(amount(order.amount_paid) or amount(order.product.price if order.product else 0) for order in orders)
    booking_revenue = sum(amount(booking.amount_paid) for booking in bookings)
    expense_total = sum(amount(expense.amount) for expense in expenses)
    revenue_total = order_revenue + booking_revenue

    days = [(datetime.utcnow().date() - timedelta(days=offset)) for offset in range(6, -1, -1)]
    income_by_day = {day: 0 for day in days}
    expense_by_day = {day: 0 for day in days}

    for order in orders:
        if order.created_at and order.created_at.date() in income_by_day:
            income_by_day[order.created_at.date()] += amount(order.amount_paid) or amount(order.product.price if order.product else 0)
    for booking in bookings:
        if booking.starts_at and booking.starts_at.date() in income_by_day:
            income_by_day[booking.starts_at.date()] += amount(booking.amount_paid)
    for expense in expenses:
        if expense.spent_at and expense.spent_at.date() in expense_by_day:
            expense_by_day[expense.spent_at.date()] += amount(expense.amount)

    return {
        "order_revenue": order_revenue,
        "booking_revenue": booking_revenue,
        "revenue_total": revenue_total,
        "expense_total": expense_total,
        "profit": revenue_total - expense_total,
        "labels": [day.strftime("%d.%m") for day in days],
        "income_series": [income_by_day[day] for day in days],
        "expense_series": [expense_by_day[day] for day in days],
    }


async def compute_finance(db: AsyncSession, club_ids: list[int]) -> dict:
    """Финансовая сводка БЕЗ загрузки всех строк (R2): гранд-тоталы через SUM,
    7-дневный график — по ограниченному окну. Память не растёт с историей."""
    days = [(datetime.utcnow().date() - timedelta(days=offset)) for offset in range(6, -1, -1)]
    labels = [day.strftime("%d.%m") for day in days]
    if not club_ids:
        return {
            "order_revenue": 0, "booking_revenue": 0, "revenue_total": 0,
            "expense_total": 0, "profit": 0, "labels": labels,
            "income_series": [0] * 7, "expense_series": [0] * 7,
        }

    order_revenue = int((await db.execute(
        select(func.coalesce(func.sum(models.Order.amount_paid), 0))
        .join(models.Product).filter(models.Product.club_id.in_(club_ids))
    )).scalar() or 0)
    booking_revenue = int((await db.execute(
        select(func.coalesce(func.sum(models.Booking.amount_paid), 0))
        .join(models.Computer).filter(models.Computer.club_id.in_(club_ids))
    )).scalar() or 0)
    expense_total = int((await db.execute(
        select(func.coalesce(func.sum(models.Expense.amount), 0))
        .filter(models.Expense.club_id.in_(club_ids))
    )).scalar() or 0)
    revenue_total = order_revenue + booking_revenue

    window_start = datetime.utcnow() - timedelta(days=7)
    income_by_day = {day: 0 for day in days}
    expense_by_day = {day: 0 for day in days}

    res_o = await db.execute(
        select(models.Order).join(models.Product)
        .filter(models.Product.club_id.in_(club_ids), models.Order.created_at >= window_start)
    )
    for order in res_o.scalars().all():
        if order.created_at and order.created_at.date() in income_by_day:
            income_by_day[order.created_at.date()] += amount(order.amount_paid)
    res_b = await db.execute(
        select(models.Booking).join(models.Computer)
        .filter(models.Computer.club_id.in_(club_ids), models.Booking.starts_at >= window_start)
    )
    for booking in res_b.scalars().all():
        if booking.starts_at and booking.starts_at.date() in income_by_day:
            income_by_day[booking.starts_at.date()] += amount(booking.amount_paid)
    res_e = await db.execute(
        select(models.Expense)
        .filter(models.Expense.club_id.in_(club_ids), models.Expense.spent_at >= window_start)
    )
    for expense in res_e.scalars().all():
        if expense.spent_at and expense.spent_at.date() in expense_by_day:
            expense_by_day[expense.spent_at.date()] += amount(expense.amount)

    return {
        "order_revenue": order_revenue, "booking_revenue": booking_revenue,
        "revenue_total": revenue_total, "expense_total": expense_total,
        "profit": revenue_total - expense_total, "labels": labels,
        "income_series": [income_by_day[day] for day in days],
        "expense_series": [expense_by_day[day] for day in days],
    }


async def computer_number_exists(db: AsyncSession, club_id: int, number: int) -> bool:
    res = await db.execute(
        select(models.Computer.id).filter(
            models.Computer.club_id == club_id,
            models.Computer.number == number,
        )
    )
    return res.scalars().first() is not None


def parse_products_file(upload: UploadFile, payload: bytes):
    filename = (upload.filename or "").lower()
    if filename.endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(cell).strip() if cell is not None else "" for cell in values[0]] if values else []
        rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        raise ValueError("Поддерживаются только CSV и XLSX файлы")

    products = []
    for row in rows:
        normalized = {str(k).strip().lower(): v for k, v in row.items() if k}
        name = normalized.get("name") or normalized.get("название") or normalized.get("товар")
        price = normalized.get("price") or normalized.get("цена")
        image_url = normalized.get("image_url") or normalized.get("картинка") or normalized.get("фото") or ""
        if not name or price in (None, ""):
            continue
        try:
            price_int = int(float(str(price).replace(",", ".")))
        except ValueError:
            continue
        if price_int < 0:
            continue
        product_name = normalize_text(str(name))
        if not product_name:
            continue
        products.append({
            "name": product_name,
            "price": price_int,
            "image_url": normalize_text(str(image_url)) if image_url else None,
        })
    return products


async def build_admin_ai_context(db: AsyncSession, club_ids: list[int]):
    clubs_count = len(club_ids)
    computers = []
    orders_count = 0
    active_bookings = 0
    if club_ids:
        res = await db.execute(select(models.Computer).filter(models.Computer.club_id.in_(club_ids)))
        computers = res.scalars().all()
        res_orders = await db.execute(
            select(func.count(models.Order.id))
            .join(models.Product)
            .filter(models.Product.club_id.in_(club_ids))
        )
        orders_count = res_orders.scalar() or 0
        res_bookings = await db.execute(
            select(func.count(models.Booking.id))
            .join(models.Computer)
            .filter(models.Computer.club_id.in_(club_ids), models.Booking.status == "active")
        )
        active_bookings = res_bookings.scalar() or 0

    free_count = len([pc for pc in computers if pc.status == "free"])
    busy_count = len([pc for pc in computers if pc.status == "busy"])
    return (
        "Проект: CyberBooking, FastAPI async + SQLAlchemy + Jinja2. "
        "Сущности: User, Club, Game, Computer, Product, Order, Booking, Notification. "
        f"Клубов владельца {clubs_count}, ПК {len(computers)}, свободно {free_count}, "
        f"занято {busy_count}, заказов {orders_count}, активных броней {active_bookings}."
    )


@router.get("", response_class=HTMLResponse)
async def admin_page(request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user or current_user.role not in CLUB_ADMIN_ROLES:
        if current_user and current_user.role == ROLE_SUPERADMIN:
            return RedirectResponse(url="/superadmin", status_code=303)
        if current_user and current_user.role == ROLE_PENDING_OWNER:
            return RedirectResponse(url="/owner/pending", status_code=303)
        return RedirectResponse(url="/login", status_code=303)

    clubs = await get_owner_clubs(db, current_user)
    club_ids = [club.id for club in clubs]

    res_games = await db.execute(select(models.Game).order_by(models.Game.name))
    all_games = res_games.scalars().all()

    orders = []
    orders_count = 0
    computers = []
    products = []
    bookings = []
    packages = []
    expenses = []
    notifications = []
    integrations = {}
    fraud_signals = []
    if club_ids:
        # R2: для таблицы показываем последние N, иначе при больших клубах грузим
        # всю историю в память. Полный счётчик — отдельным COUNT.
        res_orders = await db.execute(
            select(models.Order)
            .options(selectinload(models.Order.user), selectinload(models.Order.product))
            .join(models.Product)
            .filter(models.Product.club_id.in_(club_ids))
            .order_by(models.Order.created_at.desc())
            .limit(settings.history_page_size)
        )
        orders = res_orders.scalars().all()
        orders_count = (await db.execute(
            select(func.count(models.Order.id))
            .join(models.Product)
            .filter(models.Product.club_id.in_(club_ids))
        )).scalar() or 0

        res_computers = await db.execute(
            select(models.Computer)
            .filter(models.Computer.club_id.in_(club_ids))
            .order_by(models.Computer.club_id, models.Computer.number)
        )
        computers = res_computers.scalars().all()

        res_products = await db.execute(
            select(models.Product)
            .options(selectinload(models.Product.club))
            .filter(models.Product.club_id.in_(club_ids))
            .order_by(models.Product.id.desc())
        )
        products = res_products.scalars().all()

        res_packages = await db.execute(
            select(models.Package)
            .options(selectinload(models.Package.club))
            .filter(models.Package.club_id.in_(club_ids))
            .order_by(models.Package.club_id, models.Package.pc_category, models.Package.price)
        )
        packages = res_packages.scalars().all()

        res_bookings = await db.execute(
            select(models.Booking)
            .options(selectinload(models.Booking.computer), selectinload(models.Booking.user))
            .join(models.Computer)
            .filter(models.Computer.club_id.in_(club_ids))
            .order_by(models.Booking.starts_at.desc())
            .limit(settings.history_page_size)
        )
        bookings = res_bookings.scalars().all()

        res_expenses = await db.execute(
            select(models.Expense)
            .filter(models.Expense.club_id.in_(club_ids))
            .order_by(models.Expense.spent_at.desc())
            .limit(50)
        )
        expenses = res_expenses.scalars().all()

        res_notifications = await db.execute(
            select(models.Notification)
            .filter(models.Notification.club_id.in_(club_ids))
            .order_by(models.Notification.created_at.desc())
            .limit(15)
        )
        notifications = res_notifications.scalars().all()

        res_integrations = await db.execute(
            select(models.ClubIntegration).filter(models.ClubIntegration.club_id.in_(club_ids))
        )
        integrations = {item.club_id: item for item in res_integrations.scalars().all()}

        res_fraud = await db.execute(
            select(models.FraudSignal)
            .filter(
                or_(
                    models.FraudSignal.club_id.in_(club_ids),
                    models.FraudSignal.actor_user_id == current_user.id,
                ),
                models.FraudSignal.status == "open",
            )
            .order_by(models.FraudSignal.created_at.desc())
            .limit(10)
        )
        fraud_signals = res_fraud.scalars().all()

    total_pcs = len(computers)
    free_pcs = len([pc for pc in computers if pc.status == "free"])
    busy_pcs = len([pc for pc in computers if pc.status == "busy"])
    reserved_pcs = len([pc for pc in computers if pc.status == "reserved"])
    analytics = {
        "clubs": len(clubs),
        "products": len(products),
        "orders": orders_count,
        "total_pcs": total_pcs,
        "free_pcs": free_pcs,
        "busy_pcs": busy_pcs,
        "reserved_pcs": reserved_pcs,
        "occupancy": round((busy_pcs + reserved_pcs) / total_pcs * 100) if total_pcs else 0,
        "unread_notifications": len([n for n in notifications if not n.is_read]),
    }
    finance = await compute_finance(db, club_ids)

    return templates.TemplateResponse("admin.html", {
        "request": request,
        "clubs": clubs,
        "all_games": all_games,
        "orders": orders,
        "computers": computers,
        "products": products,
        "packages": packages,
        "bookings": bookings,
        "expenses": expenses,
        "notifications": notifications,
        "fraud_signals": fraud_signals,
        "integrations": integrations,
        "analytics": analytics,
        "finance": finance,
        "owner": current_user,
    })


@router.post("/add_club")
async def add_club(
    request: Request,
    name: str = Form(...),
    address: str = Form(...),
    city: str = Form(""),
    photo_url: str = Form(None),
    description: str = Form(""),
    amenities: str = Form(""),
    contact_phone: str = Form(None),
    working_hours: str = Form("24/7"),
    game_ids: list[int] = Form([]),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or current_user.role not in CLUB_ADMIN_ROLES:
        return RedirectResponse(url="/login", status_code=303)

    club_name = normalize_text(name)
    club_address = normalize_text(address)
    if not club_name or not club_address:
        return HTMLResponse("Название и адрес клуба обязательны", status_code=400)

    club = models.Club(
        name=club_name,
        address=club_address,
        city=normalize_text(city),
        photo_url=normalize_text(photo_url) or None,
        description=description.strip() if description else "",
        amenities=amenities.strip() if amenities else "",
        contact_phone=normalize_text(contact_phone) or None,
        working_hours=normalize_text(working_hours) or "24/7",
        owner_id=current_user.id,
        status="pending",
    )
    for gid in game_ids:
        res = await db.execute(select(models.Game).filter(models.Game.id == gid))
        game = res.scalars().first()
        if game:
            club.games.append(game)
    db.add(club)
    await db.commit()
    return RedirectResponse(url="/owner/pending", status_code=303)


@router.post("/add_computers")
async def add_computers(
    request: Request,
    club_id: int = Form(...),
    count: int = Form(...),
    category: str = Form(...),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    if count < 1 or count > 200:
        return HTMLResponse("Количество ПК должно быть от 1 до 200", status_code=400)
    if category not in VALID_PC_CATEGORIES:
        return HTMLResponse("Неизвестная категория ПК", status_code=400)

    res_max_number = await db.execute(select(func.max(models.Computer.number)).filter(models.Computer.club_id == club_id))
    last_number = res_max_number.scalar() or 0
    for i in range(1, count + 1):
        db.add(models.Computer(number=last_number + i, category=category, club_id=club_id, status="free"))
    await create_notification(db, "computer", "ПК добавлены", f"Добавлено {count} ПК категории {category}.", club_id)
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/add_computer")
async def add_computer(
    request: Request,
    club_id: int = Form(...),
    number: int = Form(...),
    category: str = Form(...),
    status: str = Form("free"),
    position_x: int = Form(0),
    position_y: int = Form(0),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    if number < 1:
        return HTMLResponse("Номер ПК должен быть положительным", status_code=400)
    if category not in VALID_PC_CATEGORIES:
        return HTMLResponse("Неизвестная категория ПК", status_code=400)
    if status not in VALID_PC_STATUSES:
        return HTMLResponse("Неизвестный статус ПК", status_code=400)
    if await computer_number_exists(db, club_id, number):
        return HTMLResponse("ПК с таким номером уже есть в этом клубе", status_code=400)

    db.add(models.Computer(
        number=number,
        category=category,
        status=status,
        club_id=club_id,
        position_x=position_x,
        position_y=position_y,
    ))
    await create_notification(db, "computer", "ПК подключен", f"ПК #{number} добавлен через админ-панель.", club_id)
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/update_computer_status")
async def update_computer_status(
    request: Request,
    pc_id: int = Form(...),
    status: str = Form(...),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    res = await db.execute(select(models.Computer).filter(models.Computer.id == pc_id))
    pc = res.scalars().first()
    if not pc:
        return HTMLResponse("ПК не найден", status_code=404)
    if not await can_manage_pc(db, current_user, pc):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    if status not in VALID_PC_STATUSES:
        return HTMLResponse("Неизвестный статус ПК", status_code=400)
    await evaluate_admin_pc_override(db, admin=current_user, pc=pc, new_status=status)
    pc.status = status
    if status == "free":
        pc.current_user_id = None
        pc.end_time = None
        await db.execute(
            update(models.Booking)
            .filter(models.Booking.computer_id == pc.id, models.Booking.status == "active")
            .values(status="completed")
        )
    await create_notification(db, "computer", "Статус ПК изменен", f"ПК #{pc.number}: {status}.", pc.club_id)
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/computers/{pc_id}/issue_ws_token")
async def issue_computer_ws_token(pc_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user:
        return JSONResponse({"detail": "Необходима авторизация"}, status_code=401)

    res = await db.execute(select(models.Computer).filter(models.Computer.id == pc_id).with_for_update())
    pc = res.scalars().first()
    if not pc:
        return JSONResponse({"detail": "ПК не найден"}, status_code=404)
    if not await can_manage_pc(db, current_user, pc):
        return JSONResponse({"detail": "Нет доступа к этому ПК"}, status_code=403)

    token = await issue_pc_token(db, pc)
    await create_notification(
        db,
        "security",
        "Выпущен токен ПК",
        f"Для ПК #{pc.number} выпущен новый WebSocket-токен.",
        pc.club_id,
    )
    await db.commit()
    return JSONResponse({
        "pc_id": pc.id,
        "token": token,
        "websocket_path": f"/ws/pc/{pc.id}",
        "header_name": "x-pc-token",
        "message": "Сохраните токен сейчас: повторно он показан не будет.",
    })


@router.post("/add_product")
async def add_product(
    request: Request,
    club_id: int = Form(...),
    name: str = Form(...),
    price: int = Form(...),
    image_url: str = Form(None),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    product_name = normalize_text(name)
    if not product_name:
        return HTMLResponse("Название товара не может быть пустым", status_code=400)
    if price < 0:
        return HTMLResponse("Цена товара не может быть отрицательной", status_code=400)
    db.add(models.Product(name=product_name, price=price, image_url=normalize_text(image_url) or None, club_id=club_id))
    await create_notification(db, "import", "Товар добавлен", f"{product_name} добавлен в магазин.", club_id)
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/add_expense")
async def add_expense(
    request: Request,
    club_id: int = Form(...),
    title: str = Form(...),
    amount_value: int = Form(..., alias="amount"),
    category: str = Form("other"),
    comment: str = Form(""),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)

    expense_title = normalize_text(title)
    expense_category = normalize_text(category) or "other"
    if not expense_title:
        return HTMLResponse("Название расхода не может быть пустым", status_code=400)
    if amount_value <= 0:
        return HTMLResponse("Сумма расхода должна быть больше нуля", status_code=400)

    db.add(models.Expense(
        club_id=club_id,
        title=expense_title,
        amount=amount_value,
        category=expense_category,
        comment=comment.strip() if comment else None,
        created_by_user_id=current_user.id,
    ))
    await create_notification(db, "finance", "Расход добавлен", f"{expense_title}: {amount_value} ₸.", club_id)
    await db.commit()
    return RedirectResponse(url="/admin#finance", status_code=303)


@router.post("/clubs/{club_id}/booking_settings")
async def save_booking_settings(
    club_id: int,
    request: Request,
    booking_mode: str = Form(...),
    booking_deposit: int = Form(0),
    cancellation_fee_percent: int = Form(0),
    db: AsyncSession = Depends(get_db),
):
    """Владелец настраивает приём броней: режим, депозит и невозвратный задаток (%)."""
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    if booking_mode not in ("request", "prepaid"):
        return HTMLResponse("Неизвестный режим бронирования", status_code=400)
    if booking_deposit < 0:
        return HTMLResponse("Депозит не может быть отрицательным", status_code=400)
    if cancellation_fee_percent < 0 or cancellation_fee_percent > 100:
        return HTMLResponse("Задаток должен быть от 0 до 100 %", status_code=400)

    res = await db.execute(select(models.Club).filter(models.Club.id == club_id))
    club = res.scalars().first()
    if not club:
        return HTMLResponse("Клуб не найден", status_code=404)
    club.booking_mode = booking_mode
    club.booking_deposit = booking_deposit if booking_mode == "prepaid" else 0
    club.cancellation_fee_percent = cancellation_fee_percent
    await db.commit()
    return RedirectResponse(url="/admin#clubs", status_code=303)


@router.post("/import_products")
async def import_products(
    request: Request,
    club_id: int = Form(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    try:
        products = parse_products_file(file, await file.read())
    except ValueError as exc:
        return HTMLResponse(str(exc), status_code=400)
    for product in products:
        db.add(models.Product(club_id=club_id, **product))
    await create_notification(db, "import", "Импорт товаров", f"Загружено товаров: {len(products)}.", club_id)
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/add_game")
async def add_game(request: Request, name: str = Form(...), db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user or current_user.role not in STAFF_ROLES:
        return RedirectResponse(url="/login", status_code=303)
    game_name = normalize_text(name)
    if not game_name:
        return HTMLResponse("Название игры не может быть пустым", status_code=400)
    res = await db.execute(select(models.Game).filter(func.lower(models.Game.name) == game_name.lower()))
    if not res.scalars().first():
        db.add(models.Game(name=game_name))
        await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/complete_order/{order_id}")
async def complete_order(order_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    res = await db.execute(
        select(models.Order)
        .options(selectinload(models.Order.product))
        .filter(models.Order.id == order_id)
    )
    order = res.scalars().first()
    if not order:
        return HTMLResponse("Заказ не найден", status_code=404)
    if not await can_manage_order(db, current_user, order):
        return HTMLResponse("Нет доступа к этому заказу", status_code=403)
    if order.status == "completed":
        return RedirectResponse(url="/admin", status_code=303)
    order.status = "completed"
    await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/cancel_order/{order_id}")
async def cancel_order_route(order_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    try:
        await OrderService(db).cancel_order(actor=current_user, order_id=order_id)
    except ServiceError as error:
        return HTMLResponse(error.message, status_code=error.status_code)
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/bookings/{booking_id}/confirm")
async def confirm_booking(booking_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    try:
        await BookingService(db).confirm_request(actor=current_user, booking_id=booking_id)
    except ServiceError as error:
        return HTMLResponse(error.message, status_code=error.status_code)
    return RedirectResponse(url="/admin#requests", status_code=303)


@router.post("/bookings/{booking_id}/reject")
async def reject_booking(booking_id: int, request: Request, db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    try:
        await BookingService(db).reject_request(actor=current_user, booking_id=booking_id)
    except ServiceError as error:
        return HTMLResponse(error.message, status_code=error.status_code)
    return RedirectResponse(url="/admin#requests", status_code=303)


@router.post("/top_up_balance")
async def top_up_balance(request: Request, user_id: int = Form(...), amount: int = Form(...), db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user or current_user.role != ROLE_SUPERADMIN:
        return HTMLResponse("Пополнение баланса доступно только суперадмину", status_code=403)
    if amount <= 0:
        return HTMLResponse("Сумма пополнения должна быть больше нуля", status_code=400)
    res = await db.execute(select(models.User).filter(models.User.id == user_id).with_for_update())
    user = res.scalars().first()
    if user:
        await credit_user_balance(
            db,
            user=user,
            actor=current_user,
            amount=amount,
            kind="admin_top_up",
            reason="Ручное пополнение баланса",
        )
        await evaluate_admin_top_up(db, admin=current_user, target_user=user, amount=amount)
        await db.commit()
    return RedirectResponse(url="/admin", status_code=303)


@router.post("/add_package")
async def add_package(
    request: Request,
    club_id: int = Form(...),
    name: str = Form(...),
    price: int = Form(...),
    paid_minutes: int = Form(...),
    bonus_minutes: int = Form(0),
    category: str = Form(...),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)
    package_name = normalize_text(name)
    if not package_name:
        return HTMLResponse("Название тарифа не может быть пустым", status_code=400)
    if price < 0:
        return HTMLResponse("Цена тарифа не может быть отрицательной", status_code=400)
    if paid_minutes < 1:
        return HTMLResponse("Оплаченные минуты должны быть больше нуля", status_code=400)
    if bonus_minutes < 0:
        return HTMLResponse("Бонусные минуты не могут быть отрицательными", status_code=400)
    duration = package_total_minutes(paid_minutes, bonus_minutes)
    if duration > MAX_PACKAGE_DURATION_MINUTES:
        return HTMLResponse("Длительность тарифа не может быть больше 24 часов", status_code=400)
    if category not in VALID_PC_CATEGORIES:
        return HTMLResponse("Неизвестная категория ПК", status_code=400)
    db.add(models.Package(
        name=package_name,
        price=price,
        duration_minutes=duration,
        paid_minutes=paid_minutes,
        bonus_minutes=bonus_minutes,
        pc_category=category,
        club_id=club_id,
    ))
    await db.commit()
    return RedirectResponse(url="/admin#packages", status_code=303)


@router.post("/integrations/1c")
async def save_1c_integration(
    request: Request,
    club_id: int = Form(...),
    base_url: str = Form(...),
    username: str = Form(""),
    secret_env_key: str = Form(""),
    health_path: str = Form("/"),
    sync_products: int = Form(0),
    sync_orders: int = Form(0),
    sync_balances: int = Form(0),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)

    integration_url = normalize_text(base_url)
    validation_error = validate_integration_url(integration_url)
    if validation_error:
        return HTMLResponse(validation_error, status_code=400)

    res = await db.execute(select(models.ClubIntegration).filter(models.ClubIntegration.club_id == club_id))
    integration = res.scalars().first()
    if not integration:
        integration = models.ClubIntegration(club_id=club_id)
        db.add(integration)

    integration.provider = "1c_http"
    integration.base_url = integration_url
    integration.username = normalize_text(username) or None
    integration.secret_env_key = normalize_text(secret_env_key) or None
    integration.health_path = normalize_text(health_path) or "/"
    integration.sync_products = 1 if sync_products else 0
    integration.sync_orders = 1 if sync_orders else 0
    integration.sync_balances = 1 if sync_balances else 0
    await db.commit()
    return RedirectResponse(url="/admin#clubs", status_code=303)


@router.post("/integrations/1c/test")
async def test_1c_integration(
    request: Request,
    club_id: int = Form(...),
    db: AsyncSession = Depends(get_db),
):
    current_user = await get_current_user(request, db)
    if not current_user or not await ensure_owner_club(db, current_user, club_id):
        return HTMLResponse("Нет доступа к этому клубу", status_code=403)

    res = await db.execute(select(models.ClubIntegration).filter(models.ClubIntegration.club_id == club_id))
    integration = res.scalars().first()
    if not integration:
        return HTMLResponse("Интеграция для клуба еще не настроена", status_code=400)

    result = await run_in_threadpool(check_integration, integration)
    touch_integration_status(integration, result.status)
    await db.commit()
    if not result.ok:
        return HTMLResponse(result.status, status_code=502)
    return RedirectResponse(url="/admin#clubs", status_code=303)


@router.post("/ai_chat")
async def admin_ai_chat(request: Request, payload: dict = Body(...), db: AsyncSession = Depends(get_db)):
    current_user = await get_current_user(request, db)
    if not current_user or current_user.role not in CLUB_ADMIN_ROLES:
        return JSONResponse({"answer": "Сначала войди как владелец клуба."}, status_code=401)

    question = (payload.get("message") or "").strip()
    if not question:
        return JSONResponse({"answer": "Напиши вопрос по проекту или админке."})

    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        return JSONResponse({
            "answer": "GROQ_API_KEY не настроен на сервере. Добавь ключ в переменные окружения и перезапусти приложение."
        }, status_code=503)

    club_ids = await owner_club_ids(db, current_user)
    body = {
        "model": os.getenv("GROQ_MODEL", "llama-3.1-8b-instant"),
        "messages": [
            {"role": "system", "content": "Ты AI-помощник администратора CyberBooking. Отвечай по-русски, кратко и практично."},
            {"role": "system", "content": await build_admin_ai_context(db, club_ids)},
            {"role": "user", "content": question},
        ],
        "temperature": 0.3,
        "max_tokens": 700,
    }
    # R1: синхронный HTTP к Groq выполняем в threadpool, чтобы НЕ блокировать
    # event loop (иначе один медленный запрос морозит все конкурентные запросы).
    answer, status_code = await run_in_threadpool(_groq_completion, api_key, body)
    return JSONResponse({"answer": answer}, status_code=status_code)


def _groq_completion(api_key: str, body: dict) -> tuple[str, int]:
    """Синхронный вызов Groq Chat Completions. Возвращает (текст_ответа|ошибка, http_status).
    Вызывать только из threadpool — внутри блокирующий urllib."""
    groq_request = urllib.request.Request(
        "https://api.groq.com/openai/v1/chat/completions",
        data=json.dumps(body).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(groq_request, timeout=30) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.read().decode("utf-8", errors="ignore"), 502
    except urllib.error.URLError as exc:
        return f"Не удалось подключиться к Groq: {exc.reason}", 502
    return data.get("choices", [{}])[0].get("message", {}).get("content", "Ответ не получен."), 200
